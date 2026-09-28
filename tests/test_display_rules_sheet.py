# -*- coding: utf-8 -*-
# Author: naohua168 <bai_bai168@qq.com>
"""展示规则 sheet 编辑层测试 — scripts/display_rules_sheet.py

覆盖 2026-09-28 重做：
  - 新结构解析（页面 / 区块 / 配置项 / 值，含 names 数组、空数组、整数、排序三选一）
  - 旧「路径」结构向前兼容
  - 校验：未知页面 / 未知配置项 / 非法排序 / 负数 → 中文报错（拒绝写回）
  - 只保留生效项：SPEC 里不再出现 部门卡 / 卡片X / 销售达成.排序 / 销售达成.销售TopN
  - 真实文件一致性：Excel「展示规则」→ dict == config/前端渲染/展示规则.json（防手改漂移）
"""
import json
from pathlib import Path

import openpyxl
import pytest

from scripts.display_rules_sheet import (
    OPTION_SHEET_NAME,
    PAGE_ORDER,
    RULE_HEADERS,
    RULE_SHEET_NAME,
    SPEC,
    build_rows,
    collect_candidates,
    excel_to_display_rules,
    parse_rows,
)

BASE_DIR = Path(__file__).parent.parent
EXCEL = BASE_DIR / "config" / "配置编辑器.xlsx"
RULES_JSON = BASE_DIR / "config" / "前端渲染" / "展示规则.json"

CUR_RULES = json.loads(RULES_JSON.read_text(encoding="utf-8"))
REMOVED = {"部门卡", "卡片1_销售达成", "卡片2_事业部矩阵", "卡片3_销售客户矩阵"}


def _sheet(rows: list[list]):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    return ws


def _roundtrip(rules: dict) -> dict:
    """JSON → sheet 行 → 解析回 dict（新结构往返）"""
    ws = _sheet([list(RULE_HEADERS)] + [list(r) for r in build_rows(rules)])
    parsed, _ = parse_rows(list(ws.iter_rows(values_only=True))[1:], list(RULE_HEADERS))
    return parsed


def _public(node) -> dict:
    """去掉 _ 前缀说明字段，便于比较配置值"""
    out = {}
    for k, v in (node or {}).items():
        if isinstance(k, str) and k.startswith("_"):
            continue
        out[k] = _public(v) if isinstance(v, dict) else v
    return out


# ══════════════════════════════════════════════════════════════
# SPEC 白名单：只留生效项
# ══════════════════════════════════════════════════════════════
class TestSpecWhitelist:
    def test_no_removed_items(self):
        assert not [s for s in SPEC if s["section"] in REMOVED], "已删除的不生效项不应再出现在 SPEC"
        assert not [s for s in SPEC if s["key"] in ("显示",)], "布尔开关已全部移除"

    def test_overview_only_topn(self):
        assert [s["key"] for s in SPEC if s["page"] == "数据总览"] == ["销售TopN"]

    def test_sales_page_has_no_items(self):
        assert [s for s in SPEC if s["page"] == "销售达成"] == []

    def test_three_done_pages_have_matrix_items(self):
        for page in ("年度达成", "月度达成", "季度达成"):
            keys = [s["key"] for s in SPEC if s["page"] == page]
            assert keys == ["最大行数", "排序", "优先展示", "客户筛选"], f"{page} 可配置项异常: {keys}"

    def test_every_spec_has_options_and_note(self):
        for s in SPEC:
            assert s["options"].strip(), f"{s['page']}/{s['key']} 缺可选值说明"
            assert s["note"].strip(), f"{s['page']}/{s['key']} 缺作用说明"


# ══════════════════════════════════════════════════════════════
# 新结构解析
# ══════════════════════════════════════════════════════════════
class TestParseNewLayout:
    def test_roundtrip_current_json(self):
        parsed = _roundtrip(CUR_RULES)
        for page in PAGE_ORDER:
            assert _public(parsed[page]) == _public(CUR_RULES[page]), f"{page} 往返不一致"

    def test_priority_names_rows(self):
        rules = {p: dict(CUR_RULES.get(p) or {}) for p in PAGE_ORDER}
        rules["年度达成"] = {"客户矩阵": {"最大行数": 0, "排序": "达成率降序",
                                      "优先展示": ["甲公司", "广汽系"], "客户筛选": []}}
        parsed = _roundtrip(rules)
        m = parsed["年度达成"]["客户矩阵"]
        assert m["优先展示"] == ["甲公司", "广汽系"]
        assert m["客户筛选"] == []
        assert m["排序"] == "达成率降序"

    def test_empty_names_list_stays_empty(self):
        parsed = _roundtrip(CUR_RULES)
        assert parsed["月度达成"]["客户矩阵"]["优先展示"] == []
        assert parsed["月度达成"]["客户矩阵"]["客户筛选"] == []

    def test_multiple_names_in_one_cell_split(self):
        rows = [list(RULE_HEADERS)] + [list(r) for r in build_rows(CUR_RULES)]
        for r in rows[1:]:
            if r[3] == "优先展示" and r[1] == "年度达成":
                r[4] = "甲公司、乙公司;丙公司"
                break
        else:
            raise AssertionError("未找到 年度达成/优先展示 行")
        parsed, _ = parse_rows(rows[1:], list(RULE_HEADERS))
        assert parsed["年度达成"]["客户矩阵"]["优先展示"][:3] == ["甲公司", "乙公司", "丙公司"]

    def test_section_row_keeps_page(self):
        """段标题行（只有页面名）→ 该页键存在，值为空 dict"""
        rows = [list(RULE_HEADERS)] + [
            ["", p, "", "", "", "", ""] for p in PAGE_ORDER
        ]
        parsed, warns = parse_rows(rows[1:], list(RULE_HEADERS))
        assert list(parsed) == PAGE_ORDER
        assert parsed["销售达成"] == {}
        assert any("缺少配置项" in w for w in warns)

    def test_forward_fill_page(self):
        """页面列留空 → 沿用上一行的页面"""
        rows = [list(RULE_HEADERS)] + [list(r) for r in build_rows(CUR_RULES)]
        target = [r for r in rows[1:] if r[1] == "年度同比" and r[3] == "排序"]
        assert target, "未找到 年度同比/排序 行"
        target[0][1] = ""            # 页面列留空
        parsed, _ = parse_rows(rows[1:], list(RULE_HEADERS))
        assert parsed["年度同比"]["排序"] == CUR_RULES["年度同比"]["排序"]
        assert parsed["年度同比"]["最大行数"] == CUR_RULES["年度同比"]["最大行数"]


# ══════════════════════════════════════════════════════════════
# 校验（拒绝脏配置进系统）
# ══════════════════════════════════════════════════════════════
class TestValidation:
    def _base(self):
        return [list(RULE_HEADERS)] + [list(r) for r in build_rows(CUR_RULES)]

    def _mutate(self, key: str, value):
        rows = self._base()
        for r in rows[1:]:
            if r[3] == key:
                r[4] = value
                return rows
        raise AssertionError(f"未找到配置项 {key}")

    def test_unknown_page(self):
        rows = self._base()
        rows.append([99, "不存在的页面", "", "销售TopN", "1", "", ""])
        with pytest.raises(ValueError, match="未知页面"):
            parse_rows(rows[1:], list(RULE_HEADERS))

    def test_unknown_item(self):
        rows = self._base()
        rows.append([99, "年度达成", "部门卡", "显示", "true", "", ""])
        with pytest.raises(ValueError, match="不存在配置项"):
            parse_rows(rows[1:], list(RULE_HEADERS))

    def test_bad_sort_value(self):
        with pytest.raises(ValueError, match="排序取值非法"):
            parse_rows(self._mutate("排序", "乱序")[1:], list(RULE_HEADERS))

    def test_negative_int(self):
        with pytest.raises(ValueError, match="不能为负数"):
            parse_rows(self._mutate("最大行数", "-3")[1:], list(RULE_HEADERS))

    def test_non_int(self):
        with pytest.raises(ValueError, match="应为整数"):
            parse_rows(self._mutate("销售TopN", "全部")[1:], list(RULE_HEADERS))

    def test_missing_page_section(self):
        """整页被删（段标题 + 其下所有行）→ 报错，避免 JSON 丢页"""
        rows = [r for r in self._base() if r[1] != "年度同比"]
        with pytest.raises(ValueError, match="缺少页面段落"):
            parse_rows(rows[1:], list(RULE_HEADERS))

    def test_missing_item_warns_not_raises(self):
        rows = [r for r in self._base() if r[3] != "客户筛选"]
        _, warns = parse_rows(rows[1:], list(RULE_HEADERS))
        assert any("缺少配置项" in w for w in warns)


# ══════════════════════════════════════════════════════════════
# 旧「路径」结构兼容
# ══════════════════════════════════════════════════════════════
class TestLegacyLayout:
    @staticmethod
    def _legacy_sheet():
        """构造旧「路径」结构（含当时的 部门卡.显示 / 卡片X.显示 / 销售TopN 等项）"""
        rows = [["页面", "路径", "值", "说明"],
                ["数据总览", "销售TopN", "10", "TopN"],
                ["年度达成", "部门卡.显示", "true", ""],
                ["年度达成", "客户矩阵.最大行数", "0", ""],
                ["年度达成", "客户矩阵.排序", "目标合计降序", ""],
                ["年度达成", "客户矩阵.优先展示", "", ""],
                ["年度达成", "客户矩阵.优先展示.1", "甲公司", ""],
                ["年度达成", "客户矩阵.优先展示.2", "乙公司", ""],
                ["年度达成", "客户矩阵.客户筛选", "", ""],
                ["月度达成", "部门卡.显示", "true", ""],
                ["季度达成", "部门卡.显示", "true", ""],
                ["销售达成", "排序", "达成率降序", ""],
                ["销售达成", "卡片1_销售达成.显示", "true", ""],
                ["销售达成", "销售TopN", "16", ""],
                ["年度同比", "最大行数", "25", ""]]
        return _sheet(rows)

    def test_legacy_parsed(self):
        parsed = excel_to_display_rules(self._legacy_sheet())
        assert parsed["数据总览"]["销售TopN"] == 10
        assert parsed["年度达成"]["部门卡"]["显示"] is True
        assert parsed["年度达成"]["客户矩阵"]["优先展示"] == ["甲公司", "乙公司"]

    def test_new_layout_detected(self):
        ws = _sheet([list(RULE_HEADERS)] + [list(r) for r in build_rows(CUR_RULES)])
        parsed = excel_to_display_rules(ws)
        assert list(parsed) == PAGE_ORDER

    def test_header_missing_raises(self):
        with pytest.raises(ValueError, match="缺少必需列"):
            excel_to_display_rules(_sheet([["页面", "值"], ["年度达成", "1"]]))


# ══════════════════════════════════════════════════════════════
# 真实配置文件一致性与产物
# ══════════════════════════════════════════════════════════════
@pytest.mark.skipif(not EXCEL.exists(), reason="配置编辑器.xlsx 不存在")
class TestRealWorkbook:
    def _ws(self):
        return openpyxl.load_workbook(EXCEL)["展示规则"]

    def test_real_sheet_matches_json(self):
        parsed = excel_to_display_rules(self._ws())
        for page in PAGE_ORDER:
            assert _public(parsed[page]) == _public(CUR_RULES[page]), (
                f"{page} 与展示规则.json 不一致；改完 Excel 请运行 "
                f"python scripts/config_excel_to_json.py 同步"
            )

    def test_real_sheet_has_no_removed_items(self):
        ws = self._ws()
        text = "\n".join(
            " ".join(str(v) for v in row if v is not None)
            for row in ws.iter_rows(values_only=True)
        )
        for bad in ("部门卡", "卡片1_销售达成", "卡片2_事业部矩阵", "卡片3_销售客户矩阵"):
            assert bad not in text, f"「展示规则」sheet 仍含已废弃项: {bad}"

    def test_option_sheet_and_dropdowns(self):
        wb = openpyxl.load_workbook(EXCEL)
        assert OPTION_SHEET_NAME in wb.sheetnames, "缺少「下拉选项」辅助 sheet"
        ws = wb["展示规则"]
        dvs = list(ws.data_validations.dataValidation)
        kinds = {dv.type for dv in dvs}
        assert {"list", "whole"} <= kinds, f"下拉/整数校验缺失: {kinds}"
        ref = [dv.formula1 for dv in dvs if dv.type == "list" and "下拉选项" in (dv.formula1 or "")]
        assert ref, "优先展示/客户筛选 的值列未引用「下拉选项」下拉"

    def test_candidates_include_parents_and_subs(self):
        cands = collect_candidates(CUR_RULES)
        kinds = {k for _, k, _ in cands}
        assert "母公司组名" in kinds and "子公司名" in kinds
        names = {n for n, _, _ in cands}
        assert "广汽系" in names and "南方韶关" in names


# ══════════════════════════════════════════════════════════════
# --init 模板生成（写到临时路径，不动真实 Excel）
# ══════════════════════════════════════════════════════════════
def test_init_template_writes_all_sheets(tmp_path, monkeypatch):
    import scripts.config_excel_to_json as cej

    out = tmp_path / "配置编辑器_test.xlsx"
    monkeypatch.setattr(cej, "EXCEL_PATH", out)
    cej.init_excel_template()

    wb = openpyxl.load_workbook(out)
    for sheet in ("时间配置", RULE_SHEET_NAME, OPTION_SHEET_NAME, "销售归属",
                  "KPI指标", "字段映射", "说明"):
        assert sheet in wb.sheetnames, f"模板缺少 sheet: {sheet}"
    # 生成的模板可直接解析回当前 JSON
    parsed = excel_to_display_rules(wb[RULE_SHEET_NAME])
    for page in PAGE_ORDER:
        assert _public(parsed[page]) == _public(CUR_RULES[page])
    # 说明页含新第 3 节 + 第 9 节
    notes = "\n".join(str(r[0]) for r in wb["说明"].iter_rows(values_only=True) if r[0])
    assert "3. 「展示规则」sheet 列说明" in notes
    assert "9. 金额单位" in notes
