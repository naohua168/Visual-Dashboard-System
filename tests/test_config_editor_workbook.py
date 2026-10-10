# -*- coding: utf-8 -*-
# Author: naohua168 <bai_bai168@qq.com>
"""配置编辑器「整本工作簿」测试 — 表格设计 + 横幅兼容解析（2026-09-28）

覆盖：
  - 统一样式：每张表第 1 行横幅、第 2 行表头、数据自第 3 行起、标签页颜色、冻结/筛选
  - 解析兼容：带横幅（表头在第 2 行）与历史文件（表头在第 1 行）都能解析
  - `--init` 模板：6 张 sheet 齐全、说明页在最后、解析回 JSON 一致
"""
import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import scripts.config_excel_to_json as cej
from scripts.sheet_style import TAB_COLORS, banner, header_row, locate_header_row

BASE_DIR = Path(__file__).parent.parent
EXCEL = BASE_DIR / "config" / "配置编辑器.xlsx"
SHEETS = ["时间配置", "展示规则", "销售归属", "KPI指标", "字段映射", "说明"]


def _norm_fmt(fmt: str | None) -> str:
    """数字格式归一：Excel 保存会把字面量转义（yyyy\\-mm\\-dd\\ hh:mm:ss ≡ yyyy-mm-dd hh:mm:ss）"""
    return (fmt or "").replace("\\", "")


def _sheet(rows, banner_text: str | None = None, headers: list[str] | None = None):
    """构造 sheet：headers 必填；banner_text 非空 → 横幅在第 1 行、表头第 2 行，
    否则按历史布局把表头写在第 1 行"""
    wb = openpyxl.Workbook()
    ws = wb.active
    if headers:
        if banner_text is not None:
            banner(ws, banner_text, len(headers))
            header_row(ws, headers, row=2)
        else:
            header_row(ws, headers, row=1)
    for r in rows:
        ws.append(r)
    return ws


def _public(node) -> dict:
    """去掉 _ 前缀说明字段，便于比较配置值（与 test_display_rules_sheet 同口径）"""
    out = {}
    for k, v in (node or {}).items():
        if isinstance(k, str) and k.startswith("_"):
            continue
        out[k] = _public(v) if isinstance(v, dict) else v
    return out


# ══════════════════════════════════════════════════════════════
# 表头行定位
# ══════════════════════════════════════════════════════════════
class TestLocateHeaderRow:
    def test_first_row(self):
        rows = [("页面", "指标", "值(万元)"), ("年度达成", "收入", 1)]
        assert locate_header_row(rows, ("页面", "指标")) == 0

    def test_after_banner(self):
        rows = [("说明横幅", None), ("页面", "指标", "值(万元)")]
        assert locate_header_row(rows, ("页面", "指标"), max_scan=4) == 1

    def test_not_found(self):
        assert locate_header_row([("甲", "乙")], ("页面", "指标")) == -1


# ══════════════════════════════════════════════════════════════
# 解析兼容：带横幅 / 无横幅
# ══════════════════════════════════════════════════════════════
class TestBannerTolerantParsing:
    TIME_HEADERS = ["配置名", "模式", "动态策略", "开始日期", "结束日期", "年份", "月份范围", "说明"]
    TIME_ROWS = [
        ["年度累计", "static", "", "2026-01-01", "2026-09-30", "", "", "数据总览/年度达成"],
        ["月度数据", "static", "", "2026-09-01", "2026-09-30", "", "", "月度达成页"],
        ["季度累计筛选", "static", "", "2026-07-01", "2026-09-30", "", "", "季度达成页"],
        ["年基线数据", "static", "", "", "", 2025, "1-9", "Phase 0"],
        ["结算模式", "常规", "", "", "", "", "", "月底结算/常规"],
    ]

    def test_time_banner_and_legacy(self):
        with_banner = _sheet(self.TIME_ROWS, "横幅", self.TIME_HEADERS)
        legacy = _sheet(self.TIME_ROWS, None, self.TIME_HEADERS)
        a = cej.excel_to_time_config(with_banner)
        b = cej.excel_to_time_config(legacy)
        assert a == b
        assert a["年度累计"]["end_date"] == "2026-09-30"
        assert a["结算模式"]["值"] == "常规"

    def test_time_chinese_labels(self):
        """2026-09-30：模式/动态策略在 Excel 里改用中文（手填日期 / 自动计算 / 最近完整月…）"""
        rows = [
            ["年度累计", "手填日期", "", "2026-01-01", "2026-09-30", "", "", "数据总览"],
            ["月度数据", "自动计算", "最近完整月", "", "", "", "", "月度达成页"],
            ["季度累计筛选", "自动计算", "最近完整季度", "", "", "", "", "季度达成页"],
            ["年基线数据", "手填日期", "", "", "", 2025, "1-9", "Phase 0"],
            ["结算模式", "常规", "", "", "", "", "", "月底结算/常规"],
        ]
        cfg = cej.excel_to_time_config(_sheet(rows, None, self.TIME_HEADERS))
        assert "_mode" not in cfg["年度累计"]        # 静态模式不写 _mode（由 resolve_time_range 兜底）
        assert cfg["年度累计"]["start_date"] == "2026-01-01"
        assert cfg["月度数据"]["_mode"] == "dynamic"
        assert cfg["月度数据"]["_strategy"] == "last_full_month"
        assert cfg["季度累计筛选"]["_strategy"] == "last_full_quarter"
        assert cfg["年基线数据"]["年份"] == 2025

    def test_time_english_legacy_still_works(self):
        """旧文件里的 static / dynamic / last_full_month 仍然能解析（向后兼容）"""
        rows = [
            ["年度累计", "static", "", "2026-01-01", "2026-09-30", "", "", ""],
            ["月度数据", "dynamic", "last_full_month", "", "", "", "", ""],
            ["季度累计筛选", "static", "", "2026-07-01", "2026-09-30", "", "", ""],
            ["年基线数据", "static", "", "", "", 2025, "1-9", ""],
            ["结算模式", "常规", "", "", "", "", "", ""],
        ]
        cfg = cej.excel_to_time_config(_sheet(rows, None, self.TIME_HEADERS))
        assert cfg["月度数据"]["_mode"] == "dynamic"
        assert cfg["月度数据"]["_strategy"] == "last_full_month"

    def test_time_illegal_mode_reports_choices(self):
        """非法模式（如手滑打成 dynami）报错要列出可选值，便于自查"""
        rows = [
            ["年度累计", "dynami", "", "2026-01-01", "2026-09-30", "", "", ""],
            ["月度数据", "手填日期", "", "2026-09-01", "2026-09-30", "", "", ""],
            ["季度累计筛选", "手填日期", "", "2026-07-01", "2026-09-30", "", "", ""],
            ["年基线数据", "手填日期", "", "", "", 2025, "1-9", ""],
            ["结算模式", "常规", "", "", "", "", "", ""],
        ]
        with pytest.raises(ValueError, match="未知模式"):
            cej.excel_to_time_config(_sheet(rows, None, self.TIME_HEADERS))

    def test_time_column_order_independent(self):
        """列顺序调整（按表头名定位）也能解析"""
        heads = ["模式", "配置名", "开始日期", "结束日期", "动态策略", "年份", "月份范围", "说明"]
        rows = [[r[1], r[0], r[3], r[4], r[2], r[5], r[6], r[7]] for r in self.TIME_ROWS]
        ws = _sheet(rows, "横幅", heads)
        cfg = cej.excel_to_time_config(ws)
        assert cfg["月度数据"]["start_date"] == "2026-09-01"

    def test_time_start_keeps_seconds(self):
        """开始日期也能精确到秒（2026-10-09）：填了时分秒就照原样写进 JSON"""
        rows = [
            ["年度累计", "手填日期", "", "2026-01-01 08:30:15", "2026-10-31 23:59:59", "", "", ""],
            ["月度数据", "手填日期", "", "2026-10-01 00:00:00", "2026-10-31 23:59:59", "", "", ""],
            ["季度累计筛选", "手填日期", "", "2026-10-01", "2026-12-31", "", "", ""],
            ["年基线数据", "手填日期", "", "", "", 2025, "1-9", ""],
            ["结算模式", "常规", "", "", "", "", "", ""],
        ]
        cfg = cej.excel_to_time_config(_sheet(rows, None, self.TIME_HEADERS))
        assert cfg["年度累计"]["start_date"] == "2026-01-01 08:30:15"
        assert cfg["月度数据"]["start_date"] == "2026-10-01 00:00:00"

    def test_month_range_date_trap_reports_hint(self):
        """月份范围被 Excel 存成日期序列号（46031）时，报错要提示把单元格设为文本"""
        rows = [
            ["年度累计", "手填日期", "", "2026-01-01", "2026-10-31", "", "", ""],
            ["月度数据", "手填日期", "", "2026-10-01", "2026-10-31", "", "", ""],
            ["季度累计筛选", "手填日期", "", "2026-10-01", "2026-12-31", "", "", ""],
            ["年基线数据", "手填日期", "", "", "", 2025, 46031, ""],
            ["结算模式", "常规", "", "", "", "", "", ""],
        ]
        with pytest.raises(ValueError, match="文本"):
            cej.excel_to_time_config(_sheet(rows, None, self.TIME_HEADERS))

    def test_kpi_banner(self):
        heads = ["页面", "指标", "值(万元)", "说明"]
        rows = [["年度达成", "收入", 50000, "覆盖"], ["年度达成", "回款", "", ""],
                ["月度达成", "收入", "", ""], ["月度达成", "回款", "", ""],
                ["季度达成", "收入", "", ""], ["季度达成", "回款", "", ""]]
        kpi = cej.excel_to_kpi(_sheet(rows, "横幅", heads))
        assert kpi["年度达成"]["收入"] == 50000.0
        assert kpi["月度达成"]["回款"] is None

    def test_attribution_banner(self):
        heads = ["母公司", "子公司", "指标", "部门", "销售", "比例", "说明"]
        rows = [["甲公司", "甲公司", "收入/回款", "检测", "张三", "1", ""],
                ["甲公司", "乙公司", "收入", "信息", "李四", "0.5", ""]]
        att = cej.excel_to_attribution(_sheet(rows, "横幅", heads))
        assert att["甲公司"]["子公司"]["乙公司"]["收入"]["信息"] == {"李四": 0.5}
        assert att["甲公司"]["子公司"]["甲公司"]["回款"]["检测"] == {"张三": 1.0}

    def test_report_dup_names_warns_only_on_symbol_diff(self, capsys):
        """仅符号/空白差异的"同一家公司写两遍" → 提示；内容不同的名字不提示"""
        cej.report_dup_names({
            "法雷奥汽车内部控制(深圳)有限公司": {
                "子公司": {"法雷奥汽车内部控制（深圳）有限公司": {}},
            },
        })
        assert "写成两遍" in capsys.readouterr().out

        cej.report_dup_names({
            "法雷奥汽车内部控制（深圳）有限公司": {
                "子公司": {"法雷奥汽车内部控制（上海）有限公司": {}},
            },
        })
        assert capsys.readouterr().out == ""

    def test_report_dup_names_ignores_multi_group(self, capsys):
        """多组配置（同一子公司在不同父组下各一份）不应提示（2026-09-29 细化规则）"""
        cej.report_dup_names({
            "广东自有客户": {"子公司": {"佛吉亚汽车部件（佛山）有限公司": {}}},
            "零部件客户": {"子公司": {"佛吉亚汽车部件(佛山)有限公司": {}}},
        })
        assert capsys.readouterr().out == ""

    def test_attribution_metric_legacy_comma(self):
        """历史文件的逗号写法「收入,回款」仍能解析（新写法 = 斜杠，2026-09-29）"""
        heads = ["母公司", "子公司", "指标", "部门", "销售", "比例", "说明"]
        rows = [["甲公司", "甲公司", "收入,回款", "检测", "张三", "1", ""]]
        att = cej.excel_to_attribution(_sheet(rows, "横幅", heads))
        assert set(att["甲公司"]["子公司"]["甲公司"]) == {"收入", "回款"}

    def test_split_metric_value(self):
        assert cej.split_metric_value("收入/回款") == ["收入", "回款"]
        assert cej.split_metric_value("收入,回款") == ["收入", "回款"]
        assert cej.split_metric_value(" 收入 ") == ["收入"]
        assert cej.split_metric_value("") == []

    def test_time_missing_header_raises(self):
        ws = _sheet([["甲", "乙"]], "横幅", ["甲", "乙"])
        with pytest.raises(ValueError, match="找不到表头行"):
            cej.excel_to_time_config(ws)


# ══════════════════════════════════════════════════════════════
# 真实工作簿：统一设计
# ══════════════════════════════════════════════════════════════
@pytest.mark.skipif(not EXCEL.exists(), reason="配置编辑器.xlsx 不存在")
class TestRealWorkbookDesign:
    def _wb(self):
        return openpyxl.load_workbook(EXCEL)

    def test_sheet_set_and_order(self):
        assert self._wb().sheetnames == SHEETS

    def test_every_sheet_has_banner(self):
        """每张表第 1 行是横幅（说明页是标题行），表头在第 2 行"""
        wb = self._wb()
        for name in SHEETS:
            ws = wb[name]
            first = ws.cell(row=1, column=1).value
            assert first, f"{name} 缺少第 1 行横幅/标题"

    def test_tab_colors_and_freeze(self):
        wb = self._wb()
        for name in SHEETS:
            ws = wb[name]
            tab = ws.sheet_properties.tabColor
            assert tab is not None, f"{name} 缺少标签页颜色"
            assert str(tab.rgb).endswith(TAB_COLORS[name])
        for name in ("时间配置", "展示规则", "销售归属", "KPI指标", "字段映射"):
            assert wb[name].freeze_panes, f"{name} 未冻结窗格"

    def test_editable_columns_highlighted(self):
        """可编辑列（时间配置的模式/日期、销售归属的销售/比例、展示规则的值）为浅黄底"""
        wb = self._wb()
        assert wb["时间配置"].cell(row=3, column=2).fill.fgColor.rgb.endswith("FFF9E6")
        assert wb["销售归属"].cell(row=3, column=5).fill.fgColor.rgb.endswith("FFF9E6")
        ws = wb["展示规则"]
        data_row = next(r for r in range(3, ws.max_row + 1)
                        if ws.cell(row=r, column=4).value == "销售TopN")
        assert ws.cell(row=data_row, column=5).fill.fgColor.rgb.endswith("FFF9E6"), "值列未高亮"

    def test_sales_sheet_banded_by_parent(self):
        """销售归属：按母公司分块斑马纹（相邻母公司底色不同）"""
        ws = self._wb()["销售归属"]
        parents, fills = [], []
        for r in range(3, ws.max_row + 1):
            p = ws.cell(row=r, column=1).value
            if p and (not parents or parents[-1] != p):
                parents.append(p)
                fills.append(ws.cell(row=r, column=1).fill.fgColor.rgb)
        assert len(parents) > 5, "母公司数异常"
        assert len(set(fills)) >= 2, "未按母公司分块着色"

    def test_autofilter_on_data_sheets(self):
        """筛选范围从表头行开始即可（Excel 保存后会把范围规范化为整块数据区）"""
        wb = self._wb()
        for name, first_col in (("时间配置", "A2"), ("销售归属", "A2"), ("字段映射", "A2")):
            ref = wb[name].auto_filter.ref or ""
            assert ref.startswith(f"{first_col}:"), f"{name} 自动筛选范围异常: {ref!r}"

    def test_time_date_cells_show_seconds(self):
        """开始/结束日期都写成日期型单元格且格式到秒（2026-10-09：开始时间也精确到秒）

        Excel 保存会给字面量加转义反斜杠（`yyyy\\-mm\\-dd\\ hh:mm:ss`），语义相同 → 比较前归一。
        """
        ws = self._wb()["时间配置"]
        for row in range(3, 6):          # 年度累计 / 月度数据 / 季度累计筛选
            for col in (4, 5):           # 开始日期 / 结束日期
                c = ws.cell(row=row, column=col)
                assert isinstance(c.value, datetime), f"r{row}c{col} 不是日期型单元格: {c.value!r}"
                assert _norm_fmt(c.number_format) == _norm_fmt(cej.DATE_CELL_FMT), \
                    f"r{row}c{col} 数字格式未到秒: {c.number_format!r}"

    def test_month_range_cell_kept_as_text(self):
        """月份范围单元格是文本格式，避免 Excel 把 1-9 存成日期 2026/1/9

        具体月份由用户维护（如 1-9 / 1-10），这里只断言「仍是文本 + 能正常解析」。
        """
        ws = self._wb()["时间配置"]
        row = next(r for r in range(3, ws.max_row + 1)
                   if ws.cell(row=r, column=1).value == "年基线数据")
        c = ws.cell(row=row, column=7)
        assert c.number_format == "@", f"月份范围单元格不是文本格式: {c.number_format!r}"
        assert re.fullmatch(r"\d{1,2}-\d{1,2}", str(c.value).strip()), \
            f"月份范围被表格软件转成了非文本值: {c.value!r}"
        months = cej.excel_to_time_config(ws)["年基线数据"]["月份范围"]
        assert len(months) == 2 and months[0] <= months[1], f"月份范围解析异常: {months}"

    def test_kpi_and_rules_keep_dropdowns(self):
        wb = self._wb()
        assert len(wb["KPI指标"].data_validations.dataValidation) == 2   # 页面 / 指标
        assert len(wb["时间配置"].data_validations.dataValidation) == 3  # 模式 / 动态策略 / 结算模式
        assert len(wb["销售归属"].data_validations.dataValidation) == 2  # 指标 / 部门
        assert len(wb["展示规则"].data_validations.dataValidation) == 3  # 排序 / 整数 / 客户名提示

    def test_metric_dropdown_offers_merged_option(self):
        """指标下拉必须能直接选到「收入/回款」（斜杠不与内联列表的逗号分隔符冲突）"""
        ws = self._wb()["销售归属"]
        formulas = [dv.formula1 or "" for dv in ws.data_validations.dataValidation
                    if dv.type == "list"]
        assert any("收入/回款" in f for f in formulas), f"指标下拉缺少合并项: {formulas}"

    def test_note_sheet_section_titles_styled(self):
        ws = self._wb()["说明"]
        texts = {r: ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)}
        sections = [r for r, v in texts.items()
                    if isinstance(v, str) and v[:2] in ("1.", "2.", "3.")]
        assert sections, "说明页缺少分节标题"
        for r in sections:
            assert ws.cell(row=r, column=1).font.bold, f"说明页第 {r} 行分节标题未加粗"


# ══════════════════════════════════════════════════════════════
# 日期容错（2026-10-10）：斜杠/中文/单位月/全角/序列号，非法日历日期按月末截断
# ══════════════════════════════════════════════════════════════
class TestDateTolerance:
    def test_common_formats(self):
        from datetime import datetime as dt
        from scripts.editor.time_config import parse_date_text

        cases = {
            "2026-09-30": dt(2026, 9, 30),
            "2026-9-30": dt(2026, 9, 30),
            "2026/9/30": dt(2026, 9, 30),
            "2026/09/30": dt(2026, 9, 30),
            "2026.9.30": dt(2026, 9, 30),
            "2026年9月30日": dt(2026, 9, 30),
            "2026年9月30": dt(2026, 9, 30),
            "２０２６／９／３０": dt(2026, 9, 30),          # 全角
            "2026/9/30 23:59": dt(2026, 9, 30, 23, 59),
            "2026/9/30 23:59:59": dt(2026, 9, 30, 23, 59, 59),
            "2026-09-30T23:59:59": dt(2026, 9, 30, 23, 59, 59),
            "2026/9/30 9:05:00": dt(2026, 9, 30, 9, 5),
            " 2026-09-30 ": dt(2026, 9, 30),
        }
        for text, expect in cases.items():
            assert parse_date_text(text) == expect, f"{text!r} 解析错误"
        assert parse_date_text(dt(2026, 9, 30, 12, 0)) == dt(2026, 9, 30, 12, 0)
        assert parse_date_text(date(2026, 9, 30)) == dt(2026, 9, 30)

    def test_excel_serial_number(self, capsys):
        from datetime import datetime as dt, timedelta
        from scripts.editor.time_config import parse_date_text

        assert parse_date_text("46031") == dt(1899, 12, 30) + timedelta(days=46031)
        assert "序列号" in capsys.readouterr().out     # 提示单元格格式该设成日期

    def test_invalid_day_clamped_to_month_end(self, capsys):
        """日志里的真实值：2026/9/31 23:59:59（9 月没有 31 号）→ 按月末截断并提示"""
        from datetime import datetime as dt
        from scripts.editor.time_config import parse_date_text

        assert parse_date_text("2026/9/31 23:59:59") == dt(2026, 9, 30, 23, 59, 59)
        out = capsys.readouterr().out
        assert "不是有效日期" in out and "2026 年 9 月没有 31 号" in out

    def test_unrecognizable_reports_supported_forms(self):
        from scripts.editor.time_config import parse_date_text

        for bad in ("下周", "2026-13-01", "2026/9/30 25:00", "2026年9月", ""):
            with pytest.raises(ValueError) as e:
                parse_date_text(bad)
            assert "日期" in str(e.value)

    def test_normalize_keeps_iso_and_fixes_others(self):
        from scripts.editor.time_config import normalize_date_text

        # 规范形态原样返回（保证不改现有配置的 JSON 文本）
        assert normalize_date_text("2026-10-01") == "2026-10-01"
        assert normalize_date_text("2026-10-31 23:59:59") == "2026-10-31 23:59:59"
        # 其它写法规范化（下游 end_of_day / pd.to_datetime 只认规范形态）
        assert normalize_date_text("2026/10/1") == "2026-10-01"
        assert normalize_date_text("2026年10月1日 08:30") == "2026-10-01 08:30:00"

    def test_time_config_accepts_slash_date(self, capsys):
        """整条解析走一遍：斜杠日期不再中断配置同步，写回 JSON 的是规范形态"""
        heads = ["配置名", "模式", "动态策略", "开始日期", "结束日期", "年份", "月份范围", "说明"]
        rows = [
            ["年度累计", "手填日期", "", "2026/9/1", "2026/9/31 23:59:59", "", "", "数据总览"],
            ["月度数据", "手填日期", "", "2026/9/1", "2026/9/30", "", "", "月度达成页"],
            ["季度累计筛选", "手填日期", "", "2026/7/1", "2026/9/30", "", "", "季度达成页"],
            ["年基线数据", "手填日期", "", "", "", "2025年", "1~8", "Phase 0"],
            ["结算模式", "常规", "", "", "", "", "", ""],
        ]
        cfg = cej.excel_to_time_config(_sheet(rows, "横幅", heads))
        assert cfg["年度累计"]["end_date"] == "2026-09-30 23:59:59"    # 月末截断 + 规范化
        assert cfg["季度累计筛选"]["start_date"] == "2026-07-01"
        assert cfg["年基线数据"]["年份"] == 2025
        assert cfg["年基线数据"]["月份范围"] == [1, 8]
        assert "不是有效日期" in capsys.readouterr().out

    def test_date_cell_value_is_tolerant(self):
        """写方向：斜杠日期也写成真正的日期单元格（不落成文本）"""
        from datetime import datetime as dt
        from scripts.editor.time_config import _date_cell_value

        assert _date_cell_value("2026/10/1") == dt(2026, 10, 1)
        assert _date_cell_value("2026-10-31 23:59:59") == dt(2026, 10, 31, 23, 59, 59)
        assert _date_cell_value("下周一") == "下周一"      # 认不出就原样返回，由读取阶段报错


# ══════════════════════════════════════════════════════════════
# --init 模板（写到临时路径）
# ══════════════════════════════════════════════════════════════
def test_init_template_design(tmp_path, monkeypatch):
    out = tmp_path / "editor_test.xlsx"
    monkeypatch.setattr(cej, "EXCEL_PATH", out)
    cej.init_excel_template()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == SHEETS, f"sheet 顺序/集合异常: {wb.sheetnames}"
    for name in SHEETS:
        assert wb[name].cell(row=1, column=1).value, f"{name} 缺少横幅"
    assert wb["展示规则"].cell(row=2, column=4).value == "配置项"
    assert wb["时间配置"].cell(row=2, column=1).value == "配置名"

    # 时间配置：开始/结束都写成日期型单元格、格式到秒；月份范围单元格为文本（2026-10-09）
    ws_t = wb["时间配置"]
    assert isinstance(ws_t.cell(row=3, column=4).value, datetime), "开始日期不是日期型单元格"
    assert ws_t.cell(row=3, column=4).number_format == cej.DATE_CELL_FMT
    assert ws_t.cell(row=3, column=5).number_format == cej.DATE_CELL_FMT
    assert ws_t.cell(row=6, column=7).number_format == "@"

    # 解析回 JSON：与事实源一致
    rules = json.loads((BASE_DIR / "config" / "前端渲染" / "展示规则.json").read_text(encoding="utf-8"))
    from scripts.display_rules_sheet import excel_to_display_rules

    parsed = excel_to_display_rules(wb["展示规则"])
    for page in rules:
        if page.startswith("_"):
            continue
        assert _public(parsed[page]) == _public(rules[page]), f"{page} 模板往返不一致"

    # 时间配置 / KPI / 销售归属 也能从带横幅的模板解析
    cfg = json.loads((BASE_DIR / "config" / "清洗配置" / "cleaning_config.json").read_text(encoding="utf-8"))
    assert cej.excel_to_time_config(wb["时间配置"])["年度累计"]["end_date"] == cfg["时间范围"]["年度累计"]["end_date"]
    assert cej.excel_to_kpi(wb["KPI指标"])["年度达成"]["收入"] is None
    att = cej.excel_to_attribution(wb["销售归属"])
    # 与 JSON 事实源同规模（不写死条数：Excel/JSON 随业务增删客户而变化）
    att_cfg = json.loads(
        (BASE_DIR / "config" / "清洗配置" / "客户销售归属.json").read_text(encoding="utf-8")
    )["客户归属"]
    assert len(att) == len(att_cfg), "模板母公司数与 JSON 不一致"
    assert sum(len(g["子公司"]) for g in att.values()) == sum(
        len(g["子公司"]) for g in att_cfg.values()), "模板子公司数与 JSON 不一致"

    # 指标列：模板统一写斜杠写法（收入/回款），不再输出逗号
    ws_a = wb["销售归属"]
    metric_vals = {ws_a.cell(row=r, column=3).value for r in range(3, ws_a.max_row + 1)}
    metric_vals.discard(None)
    assert metric_vals, "模板销售归属没有数据行"
    assert all("," not in str(v) for v in metric_vals), f"模板仍输出逗号指标: {metric_vals}"
    assert any("/" in str(v) for v in metric_vals), f"模板未使用斜杠指标: {metric_vals}"


def test_blank_never_written_as_text(tmp_path, monkeypatch):
    """空值只能写成真空单元格，不能写成字符串 "None"（2026-10-10 审计固化）

    起因：核对"生成版工作簿"与"人工编辑版"时发现 2561 格不同，一度怀疑空值被写成 "None"。
    逐格核对结论：差异全部是「人工编辑版（Excel 保存过）里是空字符串、生成版是真空单元格」——
    两者在 Excel 里显示一致，且解析结果逐字节一致（不构成缺陷）。本测试把"不许写 None 文本"钉住。
    """
    out = tmp_path / "editor_blank.xlsx"
    monkeypatch.setattr(cej, "EXCEL_PATH", out)
    cej.init_excel_template()

    bad = []
    for ws in openpyxl.load_workbook(out).worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ("None", "nan", "NaN", "NaT"):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value!r}")
    assert not bad, f"模板把空值写成了文本: {bad[:8]}"

    # 空字符串与真空单元格必须解析成同一个结果（Excel 存过的文件里两种都会出现）
    kpi_heads = ["页面", "指标", "值(万元)", "说明"]
    as_empty = [["年度达成", "收入", "", ""], ["月度达成", "收入", "", ""]]
    as_none = [["年度达成", "收入", None, None], ["月度达成", "收入", None, None]]
    assert cej.excel_to_kpi(_sheet(as_empty, "横幅", kpi_heads)) == \
        cej.excel_to_kpi(_sheet(as_none, "横幅", kpi_heads))
    assert cej.excel_to_kpi(_sheet(as_empty, "横幅", kpi_heads))["月度达成"]["收入"] is None

    time_heads = ["配置名", "模式", "动态策略", "开始日期", "结束日期", "年份", "月份范围", "说明"]
    rows_T = [("年度累计", "2026-01-01", "2026-09-30"),
              ("月度数据", "2026-10-01", "2026-10-31"),
              ("季度累计筛选", "2026-10-01", "2026-12-31")]
    t_empty = [[k, "手填日期", "", s, e, "", "", ""] for k, s, e in rows_T]
    t_empty.append(["年基线数据", "手填日期", "", "", "", 2025, "1-9", ""])
    t_none = [[k, "手填日期", None, s, e, None, None, None] for k, s, e in rows_T]
    t_none.append(["年基线数据", "手填日期", None, None, None, 2025, "1-9", None])
    assert cej.excel_to_time_config(_sheet(t_empty, "横幅", time_heads)) == \
        cej.excel_to_time_config(_sheet(t_none, "横幅", time_heads))


def test_init_template_is_byte_stable(tmp_path, monkeypatch):
    """同一份配置两次生成，xlsx **逐字节一致**（2026-10-10）

    做法：固定文档属性（作者/创建与修改时间）+ 规范化 zip（条目排序、固定时间戳、
    钉死 docProps/core.xml 里被 openpyxl 强制改写的 dcterms:modified）。
    作用：可以拿 sha256 判断"生成版有没有被人手改过 / 生成器有没有意外改动设计"。
    """
    outs = [tmp_path / "gen1.xlsx", tmp_path / "gen2.xlsx"]
    for out in outs:
        monkeypatch.setattr(cej, "EXCEL_PATH", out)
        cej.init_excel_template()
    b1, b2 = outs[0].read_bytes(), outs[1].read_bytes()
    assert len(b1) == len(b2), f"两次生成大小不同: {len(b1)} vs {len(b2)}"
    assert b1 == b2, "两次生成字节不同（生成不确定：检查文档属性/zip 时间戳是否被固定）"

    # 生成版里不允许出现空字符串单元格（Excel 存盘会留下这种残留，生成器统一写真空）
    wb = openpyxl.load_workbook(outs[0])
    blanks = [f"{ws.title}!{c.coordinate}"
              for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value == ""]
    assert not blanks, f"生成的工作簿存在空字符串单元格: {blanks[:8]}"


def test_canonicalize_workbook_inplace(tmp_path):
    """--normalize：清空字符串单元格 + 内容不动 + 重复规范化字节不变

    注意：openpyxl 赋 `""` 等于赋空（单元格都不会建），所以"空字符串单元格"只能在
    已读出的工作簿对象上直接塞（Excel 存盘的共享字符串读出来正是 `''`）。
    """
    from scripts.sheet_style import canonicalize_workbook, clear_blank_strings, new_workbook

    # ① 空字符串清理：只动值，其他内容原地保留
    wb = new_workbook()
    ws = wb.active
    ws["A1"] = "客户"
    ws["A2"] = "上海某某有限公司"
    blank = ws.cell(row=3, column=4)
    blank._value = ""                       # 模拟 Excel 存盘读出来的空字符串
    assert clear_blank_strings(wb) == 1
    assert ws["D3"].value is None
    assert ws["A2"].value == "上海某某有限公司"

    # ② 原地规范化：无空串时报告 0，且重复执行字节不变（幂等）
    from scripts.sheet_style import save_workbook

    path = tmp_path / "edit.xlsx"
    save_workbook(wb, path)
    stat = canonicalize_workbook(path)
    assert stat["blank_cleared"] == 0
    assert openpyxl.load_workbook(path).active["A2"].value == "上海某某有限公司"

    once = path.read_bytes()
    canonicalize_workbook(path)
    assert path.read_bytes() == once, "重复规范化会改动文件字节"


def _make_note_workbook(path: Path, head: list[str]) -> None:
    """造一个只有「说明」sheet 的临时工作簿（模拟"旧版说明书"）"""
    wb = openpyxl.Workbook()
    wb.active.title = "说明"
    ws = wb.active
    ws.cell(row=1, column=1).value = "配置编辑器使用说明"
    for i, text in enumerate(head, start=2):
        ws.cell(row=i, column=1).value = text
    wb.save(path)


def test_refresh_note_sheet_appends_and_rewrites(tmp_path, monkeypatch):
    """--refresh-notes：缺的分节追加到末尾；已有但文案变了的按分节**原地重写**（其它行不动）"""
    out = tmp_path / "editor_copy.xlsx"
    head = ["1. 本 Excel 是「编辑层」", cej.UNIT_NOTE_LINES[0][0],
            "   - 旧文案：这句应被同分节的规范文案替换"]
    _make_note_workbook(out, head)
    monkeypatch.setattr(cej, "EXCEL_PATH", out)

    cej.refresh_note_sheet()

    ws = openpyxl.load_workbook(out)["说明"]
    texts = [str(ws.cell(row=r, column=1).value).strip() for r in range(1, ws.max_row + 1)]
    # 追加：第 10/11/12 节都在
    assert any(t.startswith("10.") and "名称匹配" in t for t in texts), "未追加第 10 节"
    assert any(t.startswith("11.") and "数据质量与产物" in t for t in texts), "未追加第 11 节"
    assert any(t.startswith("12.") and "业务规则.json" in t for t in texts), "未追加第 12 节"
    # 更新：第 9 节被写回规范文案，旧文案消失（口径改了要能纠正旧说明书）
    assert all("这句应被同分节的规范文案替换" not in t for t in texts), "旧文案未被替换"
    assert str(cej.UNIT_NOTE_LINES[-1][0]).strip() in texts, "规范文案未写入"
    # 结构：标题 + 第 1 节 + 规范化后的第 9 节 + 3 个新分节（每节前空一行）
    assert [ws.cell(row=r, column=1).value for r in (1, 3)] == ["配置编辑器使用说明", head[1]]
    assert ws.max_row == 2 + len(cej.UNIT_NOTE_LINES) + sum(
        1 + len(b) for b in (cej.NAME_NOTE_LINES, cej.OUTPUT_NOTE_LINES, cej.BIZ_NOTE_LINES))
    for block in (cej.UNIT_NOTE_LINES, cej.NAME_NOTE_LINES, cej.OUTPUT_NOTE_LINES, cej.BIZ_NOTE_LINES):
        assert texts.count(str(block[0][0]).strip()) == 1, f"分节重复: {block[0][0]}"

    # 再次执行：内容与行数都不变（幂等）
    cej.refresh_note_sheet()
    ws2 = openpyxl.load_workbook(out)["说明"]
    assert ws2.max_row == ws.max_row
    assert [str(ws2.cell(row=r, column=1).value).strip()
            for r in range(1, ws2.max_row + 1)] == texts
