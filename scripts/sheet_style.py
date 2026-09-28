# Author: naohua168 <bai_bai168@qq.com>
"""配置编辑器 — 统一的表格视觉规范（2026-09-28）

所有 sheet 共用同一套配色与结构，避免各页风格不一：

    ┌─ 横幅（第 1 行，合并整行）：这张表干什么、只改哪一列
    ├─ 表头（第 2 行，深蓝底白字加粗居中）
    └─ 数据区（第 3 行起）
         · 可编辑列 → 浅黄底 (FFF9E6)
         · 只读/参考列 → 浅灰底 (F7F7F7)
         · 分组斑马纹 → 每 N 行浅蓝底 (F7FBFF)，用于「销售归属」按母公司分组

同时提供 `locate_header_row()`：解析器用它在开头若干行里定位真正的表头行，
从而兼容「有横幅」与「无横幅（历史文件）」两种布局。
"""
from __future__ import annotations

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    raise SystemExit("缺少 openpyxl，请先安装: pip install openpyxl")

# ── 配色（与看板主色一致）──
COLOR_PRIMARY = "1F4E78"        # 表头深蓝
COLOR_BANNER_BG = "EAF3FF"      # 横幅浅蓝
COLOR_SECTION_BG = "F2F7FF"     # 区块/段标题浅蓝
COLOR_EDITABLE_BG = "FFF9E6"    # 可编辑列（浅黄）
COLOR_READONLY_BG = "F7F7F7"    # 只读/参考列（浅灰）
COLOR_BAND_BG = "F7FBFF"        # 斑马纹（浅蓝）
COLOR_BORDER = "B0B7C3"
COLOR_SECTION_TITLE = "1F4E78"

# 每个 sheet 的标签页颜色（一眼区分）
TAB_COLORS = {
    "时间配置": "1F4E78",
    "展示规则": "7030A0",
    "销售归属": "2E7D32",
    "KPI指标": "C55A11",
    "字段映射": "00838F",
    "说明": "808080",
}


def thin_border() -> Border:
    side = Side(style="thin", color=COLOR_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def locate_header_row(rows: list[tuple], required: tuple[str, ...],
                      max_scan: int = 4) -> int:
    """在前 max_scan 行里找出「同时包含 required 全部列名」的行号（0-based）

    Returns:
        表头所在行下标；找不到返回 -1（调用方给出中文报错）
    """
    for i, row in enumerate(rows[:max_scan]):
        cells = {str(v).strip() for v in row if v is not None}
        if all(c in cells for c in required):
            return i
    return -1


def banner(ws, text: str, ncols: int, row: int = 1, height: int = 26) -> None:
    """第 1 行横幅：这张表怎么用（合并整行，浅蓝底 + 深蓝粗体）"""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = PatternFill("solid", fgColor=COLOR_BANNER_BG)
    c.font = Font(bold=True, size=11, color=COLOR_PRIMARY)
    c.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[row].height = height


def header_row(ws, headers: list[str], row: int = 2, height: int = 22) -> None:
    """表头行：深蓝底 / 白字 / 加粗 / 居中 / 边框"""
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = PatternFill("solid", fgColor=COLOR_PRIMARY)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border()
    ws.row_dimensions[row].height = height


def section_row(ws, row: int, text: str, ncols: int,
                first_col: int = 1) -> None:
    """区块/段标题行：浅蓝底 + 深蓝粗体（合并到末列）"""
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=ncols)
    for col in range(first_col, ncols + 1):
        ws.cell(row=row, column=col).border = thin_border()
    c = ws.cell(row=row, column=first_col, value=text)
    c.fill = PatternFill("solid", fgColor=COLOR_SECTION_BG)
    c.font = Font(bold=True, color=COLOR_SECTION_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[row].height = 20


def data_cell(ws, row: int, col: int, value, editable: bool = False,
              readonly: bool = False, wrap: bool = False,
              center: bool = False, band: bool = False):
    """数据单元格：统一边框 + 可选底色（可编辑浅黄 / 只读浅灰 / 斑马纹浅蓝）"""
    c = ws.cell(row=row, column=col, value=value)
    c.border = thin_border()
    c.alignment = Alignment(vertical="center", wrap_text=wrap,
                            horizontal="center" if center else "left")
    if editable:
        c.fill = PatternFill("solid", fgColor=COLOR_EDITABLE_BG)
    elif readonly:
        c.fill = PatternFill("solid", fgColor=COLOR_READONLY_BG)
    elif band:
        c.fill = PatternFill("solid", fgColor=COLOR_BAND_BG)
    return c


def set_widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def finish(ws, freeze: str | None = None, autofilter: str | None = None,
           tab_color: str | None = None) -> None:
    """收尾：冻结窗格 / 自动筛选 / 标签页颜色"""
    if freeze:
        ws.freeze_panes = freeze
    if autofilter:
        ws.auto_filter.ref = autofilter
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
