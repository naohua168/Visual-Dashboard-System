# Author: naohua168 <bai_bai168@qq.com>
"""命令行入口 —— 参数解析与各步骤编排

由 scripts/config_excel_to_json.py 拆分而来（行为不变，见 logs/_dryrun_before.txt 比对）。
"""
from __future__ import annotations

import json
import sys
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("缺少 openpyxl，请先安装: pip install openpyxl")
    sys.exit(1)
try:  # 直接以脚本方式运行时（python scripts/config_excel_to_json.py）
    from column_mapping_sheet import (
        MAP_SHEET_NAME, attach_mapping_sheet, excel_to_column_mapping,
        read_latest_hits, refresh_mapping_sheet, update_column_mapping,
    )
except ImportError:  # 以包形式导入时（如单元测试）
    from scripts.column_mapping_sheet import (  # type: ignore[no-redef]
        MAP_SHEET_NAME, attach_mapping_sheet, excel_to_column_mapping,
        read_latest_hits, refresh_mapping_sheet, update_column_mapping,
    )
try:  # 直接以脚本方式运行时
    from display_rules_sheet import (
        RULE_NOTE_LINES, RULE_SHEET_NAME, attach_rules_sheet,
        excel_to_display_rules, refresh_rules_sheet,
    )
except ImportError:  # 以包形式导入时（如单元测试）
    from scripts.display_rules_sheet import (  # type: ignore[no-redef]
        RULE_NOTE_LINES, RULE_SHEET_NAME, attach_rules_sheet,
        excel_to_display_rules, refresh_rules_sheet,
    )
from .attribution import excel_to_attribution, update_attribution
from .constants import BASE_DIR, EXCEL_PATH, CLEANING_CFG
from .display_rules import update_display_rules, update_kpi_rules, excel_to_kpi
from .paths import BASE_DIR, EXCEL_PATH, CLEANING_CFG
from .template import init_excel_template, refresh_note_sheet
from .time_config import excel_to_time_config, update_cleaning_config


def main():
    args = sys.argv[1:]
    init = "--init" in args
    init_map = "--init-map" in args
    init_rules = "--init-rules" in args
    refresh_notes = "--refresh-notes" in args
    dry_run = "--dry-run" in args

    if init:
        init_excel_template()
        return 0

    if refresh_notes:
        path = refresh_note_sheet()
        print(f"✅ 已刷新「说明」sheet 新增分节（其余 sheet 未改动）: {path.relative_to(BASE_DIR)}")
        return 0

    if init_map:
        path = refresh_mapping_sheet()
        print(f"✅ 已刷新「{MAP_SHEET_NAME}」sheet（其余 sheet 未改动）: {path.relative_to(BASE_DIR)}")
        return 0

    if init_rules:
        path = refresh_rules_sheet()
        print(f"✅ 已刷新「{RULE_SHEET_NAME}」sheet（其余 sheet 未改动）: "
              f"{path.relative_to(BASE_DIR)}")
        return 0

    if not EXCEL_PATH.exists():
        print(f"未找到配置编辑器 Excel: {EXCEL_PATH.relative_to(BASE_DIR)}")
        print("请先运行: python scripts/config_excel_to_json.py --init")
        return 1

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # 1) 时间配置 → cleaning_config.json
    if "时间配置" in wb.sheetnames:
        ws = wb["时间配置"]
        time_config = excel_to_time_config(ws)
        new_cfg = update_cleaning_config(time_config, dry_run=dry_run)
        if dry_run:
            print("【DRY-RUN】将写回以下时间范围：")
            print(json.dumps(new_cfg["时间范围"], ensure_ascii=False, indent=2))
        else:
            print("✅ 已更新 cleaning_config.json 的「时间范围」：")
            for k, v in time_config.items():
                print(f"  - {k}: {json.dumps(v, ensure_ascii=False)}")
    else:
        print("⚠️  Excel 缺少「时间配置」sheet，跳过")

    # 2) 展示规则 → 展示规则.json
    if RULE_SHEET_NAME in wb.sheetnames:
        ws_r = wb[RULE_SHEET_NAME]
        rules = excel_to_display_rules(ws_r)
        new_rules = update_display_rules(rules, dry_run=dry_run)
        if dry_run:
            print("【DRY-RUN】将写回以下展示规则：")
            print(json.dumps(new_rules, ensure_ascii=False, indent=2))
        else:
            print("✅ 已更新 展示规则.json：")
            for page, obj in rules.items():
                keys = list(obj.keys())
                print(f"  - {page}: {keys}")
    else:
        print("⚠️  Excel 缺少「展示规则」sheet，跳过")

    # 3) KPI指标 → 展示规则.json 各页面 KPI指标 区块
    if "KPI指标" in wb.sheetnames:
        ws_k = wb["KPI指标"]
        kpi = excel_to_kpi(ws_k)
        new_rules = update_kpi_rules(kpi, dry_run=dry_run)
        if dry_run:
            print("【DRY-RUN】将写回以下 KPI 指标覆盖：")
            for page, m in kpi.items():
                print(f"  - {page}: {json.dumps(m, ensure_ascii=False)}")
        else:
            set_count = sum(
                1 for m in kpi.values() for v in m.values() if v is not None
            )
            print(f"✅ 已更新 展示规则.json 的 KPI指标 区块（{set_count} 项覆盖，其余按指标文件合计）")
    else:
        print("⚠️  Excel 缺少「KPI指标」sheet，跳过")

    # 4) 销售归属 → 客户销售归属.json
    if "销售归属" in wb.sheetnames:
        ws_a = wb["销售归属"]
        att = excel_to_attribution(ws_a)
        new_att = update_attribution(att, dry_run=dry_run)
        if dry_run:
            print("【DRY-RUN】将写回以下销售归属：")
            print(f"  母公司数: {len(new_att['客户归属'])}")
        else:
            n_subs = sum(len(g.get("子公司", {})) for g in att.values())
            print(f"✅ 已更新 客户销售归属.json：{len(att)} 母公司 / {n_subs} 子公司")
    else:
        print("⚠️  Excel 缺少「销售归属」sheet，跳过")

    # 5) 字段映射 → cleaning_config.json 的 数据源.*.列映射
    if MAP_SHEET_NAME in wb.sheetnames:
        cfg_now = json.loads(CLEANING_CFG.read_text(encoding="utf-8"))
        mapping, warns = excel_to_column_mapping(wb[MAP_SHEET_NAME], cfg_now)
        for w in warns:
            print(f"  ⚠️ {w}")
        _, changed = update_column_mapping(mapping, dry_run=dry_run)
        if dry_run:
            print("【DRY-RUN】列映射将发生变更：")
            for c in (changed or ["（无变更）"]):
                print(f"  - {c}")
        elif changed:
            print(f"✅ 已更新 cleaning_config.json 的「列映射」：{len(changed)} 个来源变更 → {', '.join(changed)}")
        else:
            print("✅ 字段映射与 JSON 一致（无需变更）")
    else:
        print(f"⚠️  Excel 缺少「{MAP_SHEET_NAME}」sheet，跳过")

    return 0
