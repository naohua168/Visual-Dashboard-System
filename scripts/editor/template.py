# Author: naohua168 <bai_bai168@qq.com>
"""--init 系列 —— 从 JSON 生成/刷新「配置编辑器.xlsx」各 sheet

由 scripts/config_excel_to_json.py 拆分而来（行为不变，见 logs/_dryrun_before.txt 比对）。
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("缺少 openpyxl，请先安装: pip install openpyxl")
    sys.exit(1)
try:  # 直接以脚本方式运行时（python scripts/config_excel_to_json.py）
    from column_mapping_sheet import (
        MAP_SHEET_NAME, attach_mapping_sheet, excel_to_column_mapping,
        read_latest_hits, refresh_mapping_sheet, update_column_mapping,
    )
except ImportError:  # 以包形式导入时（如单元测试）
    from scripts.column_mapping_sheet import (  # type: ignore[no-redef]
        MAP_SHEET_NAME, attach_mapping_sheet, excel_to_column_mapping,
        read_latest_hits, refresh_mapping_sheet, update_column_mapping,
    )
try:  # 直接以脚本方式运行时
    from display_rules_sheet import (
        RULE_NOTE_LINES, RULE_SHEET_NAME, attach_rules_sheet,
        excel_to_display_rules, refresh_rules_sheet,
    )
except ImportError:  # 以包形式导入时（如单元测试）
    from scripts.display_rules_sheet import (  # type: ignore[no-redef]
        RULE_NOTE_LINES, RULE_SHEET_NAME, attach_rules_sheet,
        excel_to_display_rules, refresh_rules_sheet,
    )
try:  # 直接以脚本方式运行时
    from sheet_style import (
        COLOR_BANNER_BG, COLOR_PRIMARY, TAB_COLORS, banner, data_cell, finish,
        header_row, locate_header_row, set_widths,
    )
except ImportError:  # 以包形式导入时（如单元测试）
    from scripts.sheet_style import (  # type: ignore[no-redef]
        COLOR_BANNER_BG, COLOR_PRIMARY, TAB_COLORS, banner, data_cell, finish,
        header_row, locate_header_row, set_widths,
    )
from .constants import TIME_KEYS, MODE_LABELS, STRATEGY_LABELS, DATE_CELL_FMT, BASE_DIR, EXCEL_PATH, CLEANING_CFG, DISPLAY_RULES_CFG, ATTRIBUTION_CFG, TIME_KEYS, MODE_LABELS, STRATEGY_LABELS, HEADERS, ATT_HEADERS, METRIC_VALUES, METRIC_BOTH, DEPT_VALUES, KPI_PAGES, KPI_METRICS, KPI_HEADERS, GROUP_NOTE_LINES, UNIT_NOTE_LINES, NAME_NOTE_LINES, OUTPUT_NOTE_LINES, BIZ_NOTE_LINES, DATE_CELL_FMT
from .paths import BASE_DIR, EXCEL_PATH, CLEANING_CFG, DISPLAY_RULES_CFG, ATTRIBUTION_CFG
from .time_config import _date_cell_value


def _cell_fill(name: str) -> str:
    fills = {
        "年度累计": "DDEBF7",
        "月度数据": "E2EFDA",
        "季度累计筛选": "FFF2CC",
        "年基线数据": "FCE4D6",
    }
    return fills.get(name, "FFFFFF")


def init_excel_template():
    """从 cleaning_config.json 当前时间范围导出 Excel 模板"""
    cfg = json.loads(CLEANING_CFG.read_text(encoding="utf-8"))
    tr = cfg.get("时间范围", {})

    wb = openpyxl.Workbook()

    # ── Sheet 1: 时间配置 ──
    ws = wb.active
    ws.title = "时间配置"
    from openpyxl.worksheet.datavalidation import DataValidation

    banner(ws, "改日期 → 保存 → 运行 启动系统.bat / run_all.bat 生效。"
               "「模式」选「手填日期」= 自己填开始/结束日期；选「自动计算」= 系统按「动态策略」自动算"
               "（最近完整月 / 最近完整季度），不用填日期；「结算模式」行的值填在「模式」列。", len(HEADERS))
    header_row(ws, HEADERS, row=2)

    row_idx = 3
    for key in TIME_KEYS:
        spec = tr.get(key, {})
        fill_color = _cell_fill(key)
        if spec.get("_mode") == "dynamic":
            values = [key, MODE_LABELS["dynamic"],
                      STRATEGY_LABELS.get(spec.get("_strategy", ""), spec.get("_strategy", "")),
                      "", "", "", "", spec.get("_使用方", "")]
        elif key == "年基线数据":
            months = spec.get("月份范围", [])
            values = [
                key, MODE_LABELS["static"], "",
                "", "", spec.get("年份", ""),
                f"{months[0]}-{months[1]}" if len(months) == 2 else "",
                spec.get("_使用方", ""),
            ]
        else:
            values = [
                key, MODE_LABELS["static"], "",
                spec.get("start_date", ""), spec.get("end_date", ""),
                "", "", spec.get("_使用方", ""),
            ]
        for col, v in enumerate(values, 1):
            if col in (4, 5):
                v = _date_cell_value(v)  # 开始/结束日期：写成 datetime，单元格显示到秒
            c = data_cell(ws, row_idx, col, v,
                          editable=col in (2, 3, 4, 5, 6, 7), readonly=(col == 8),
                          wrap=(col == 8))
            if col in (4, 5) and isinstance(c.value, datetime):
                c.number_format = DATE_CELL_FMT
            elif col == 7:  # 月份范围：文本格式，防止 Excel 把 1-9 存成日期（2026/1/9）
                c.number_format = "@"
            if col == 1:  # 配置名列保留按行的身份色（年度/月度/季度/年基线）
                c.fill = PatternFill("solid", fgColor=fill_color)
        row_idx += 1

    last_time_row = row_idx - 1   # 4 个配置行的最后一行（结算模式行不参与下面两个下拉）

    # 结算模式行（值填在「模式」列：月底结算 / 常规）
    settle = tr.get("结算模式", {})
    settle_row = None
    if settle:
        settle_row = row_idx
        values = [
            "结算模式", settle.get("值", "常规"), "",
            "", "", "", "",
            settle.get("_说明", "月底结算: 当年/季度累计=仅运营端; 常规: =财务端+运营端"),
        ]
        for col, v in enumerate(values, 1):
            c = data_cell(ws, row_idx, col, v, editable=(col == 2), readonly=(col == 8),
                          wrap=(col == 8))
            if col == 1:
                c.fill = PatternFill("solid", fgColor="FCE4D6")
        row_idx += 1

    set_widths(ws, [16, 12, 20, 18, 26, 8, 12, 46])

    dv_mode = DataValidation(type="list", formula1='"' + ",".join(MODE_LABELS.values()) + '"',
                             allow_blank=True, promptTitle="怎么填（二选一）",
                             prompt="手填日期 = 自己填「开始日期 / 结束日期」；"
                                    "自动计算 = 按右侧「动态策略」自动算，不用填日期")
    dv_mode.add(f"B3:B{last_time_row}")
    ws.add_data_validation(dv_mode)
    dv_strat = DataValidation(type="list", formula1='"' + ",".join(STRATEGY_LABELS.values()) + '"',
                              allow_blank=True, promptTitle="怎么填（仅自动计算时用）",
                              prompt="最近完整月 / 最近完整季度（例：9 月跑批 → 最近完整月 = 8 月）")
    dv_strat.add(f"C3:C{last_time_row}")
    ws.add_data_validation(dv_strat)
    if settle_row:  # 结算模式行单独给下拉（原先被普通模式下拉误罩，选不到合法值）
        dv_settle = DataValidation(type="list", formula1='"月底结算,常规"', allow_blank=True,
                                   promptTitle="怎么填", prompt="月底结算 / 常规")
        dv_settle.add(f"B{settle_row}")
        ws.add_data_validation(dv_settle)
    finish(ws, freeze="A3", autofilter="A2:H2", tab_color=TAB_COLORS["时间配置"])

    # ── Sheet 2: 展示规则 ──
    # （2026-09-28 重做：序号/页面/区块/配置项/值/可选值/说明 + 下拉验证，见 display_rules_sheet.py）
    rules = json.loads(DISPLAY_RULES_CFG.read_text(encoding="utf-8"))
    attach_rules_sheet(wb, rules)

    # ── Sheet 3: 销售归属 ──
    att = json.loads(ATTRIBUTION_CFG.read_text(encoding="utf-8"))["客户归属"]
    ws_a = wb.create_sheet("销售归属")
    banner(ws_a, "母公司 → 子公司 → 指标 → 部门 → 销售 → 比例（通常只改「销售」「比例」两列）。"
                 "一家母公司一块浅蓝斑马纹；比例 = 全额填 1，拆分填 0.3 或 30%（同一子公司各销售比例合计 = 1）。",
           len(ATT_HEADERS))
    header_row(ws_a, ATT_HEADERS, row=2)

    att_row = 3
    for group_i, (parent, group) in enumerate(att.items()):
        band = (group_i % 2 == 1)
        for sub, sub_data in group.get("子公司", {}).items():
            # 判断收入/回款是否同构：同构 → 指标=收入/回款 一行
            inc = sub_data.get("收入", {})
            pay = sub_data.get("回款", {})
            if inc == pay and inc:
                note = sub_data.get("_说明", "")
                for dept, sales_map in inc.items():
                    for sales, ratio in sales_map.items():
                        vals = [parent, sub, METRIC_BOTH, dept, sales,
                                f"{ratio:.15g}", note]
                        for c, v in enumerate(vals, 1):
                            data_cell(ws_a, att_row, c, v,
                                      editable=c in (5, 6), band=band, wrap=(c == 7))
                        att_row += 1
            else:
                # 结构不同：收入/回款分开行
                for metric, m in (("收入", inc), ("回款", pay)):
                    if not m:
                        continue
                    note = sub_data.get("_说明", "")
                    for dept, sales_map in m.items():
                        for sales, ratio in sales_map.items():
                            vals = [parent, sub, metric, dept, sales,
                                    f"{ratio:.15g}", note]
                            for c, v in enumerate(vals, 1):
                                data_cell(ws_a, att_row, c, v,
                                          editable=c in (5, 6), band=band, wrap=(c == 7))
                            att_row += 1

    set_widths(ws_a, [26, 36, 14, 10, 12, 10, 30])

    # 数据验证：指标 / 部门（取值来自文件顶部 METRIC_VALUES / DEPT_VALUES，避免重复硬编码）
    from openpyxl.worksheet.datavalidation import DataValidation
    dv_metric = DataValidation(type="list", formula1='"' + ",".join(METRIC_VALUES) + '"',
                               allow_blank=True, promptTitle="怎么填",
                               prompt="收入 / 回款 / 收入/回款（收入与回款归属一致时填「收入/回款」一行搞定）")
    dv_metric.add(f"C3:C{att_row + 50}")
    ws_a.add_data_validation(dv_metric)
    dv_dept = DataValidation(type="list", formula1='"' + ",".join(DEPT_VALUES) + '"',
                             allow_blank=True, promptTitle="怎么填",
                             prompt="检测 / 信息 / 能源 / 海外")
    dv_dept.add(f"D3:D{att_row + 50}")
    ws_a.add_data_validation(dv_dept)

    finish(ws_a, freeze="C3", autofilter="A2:G2", tab_color=TAB_COLORS["销售归属"])

    # ── Sheet 4: KPI指标（Hero 圆环指标总数覆盖，万元）──
    rules_now = json.loads(DISPLAY_RULES_CFG.read_text(encoding="utf-8"))
    ws_k = wb.create_sheet("KPI指标")
    banner(ws_k, "「值(万元)」留空 = 不覆盖（Hero 圆环按指标文件合计）；填数字 = 覆盖圆环上的指标总数。"
                 "只影响 Hero 圆环，不改表格/矩阵里的明细指标。", len(KPI_HEADERS))
    header_row(ws_k, KPI_HEADERS, row=2)

    k_row = 3
    k_notes = {
        "年度达成": "年度达成页 + 数据总览页的年度 KPI 指标总数",
        "月度达成": "月度达成页的 KPI 指标总数",
        "季度达成": "季度达成页的 KPI 指标总数",
    }
    for page in KPI_PAGES:
        page_kpi = rules_now.get(page, {}).get("KPI指标", {})
        band = (KPI_PAGES.index(page) % 2 == 1)
        for metric in KPI_METRICS:
            v = page_kpi.get(metric)
            vals = [page, metric, v if v is not None else "", k_notes[page]]
            for c, vv in enumerate(vals, 1):
                data_cell(ws_k, k_row, c, vv, editable=(c == 3), band=band, wrap=(c == 4))
            k_row += 1

    set_widths(ws_k, [16, 12, 16, 52])
    # 数据验证：页面 / 指标 下拉
    from openpyxl.worksheet.datavalidation import DataValidation
    dv_kpage = DataValidation(type="list", formula1='"' + ",".join(KPI_PAGES) + '"', allow_blank=True)
    dv_kpage.add(f"A3:A{ws_k.max_row + 50}")
    ws_k.add_data_validation(dv_kpage)
    dv_kmetric = DataValidation(type="list", formula1='"' + ",".join(KPI_METRICS) + '"', allow_blank=True)
    dv_kmetric.add(f"B3:B{ws_k.max_row + 50}")
    ws_k.add_data_validation(dv_kmetric)
    finish(ws_k, freeze="A3", tab_color=TAB_COLORS["KPI指标"])

    # ── Sheet 5: 说明 ──
    ws2 = wb.create_sheet("说明")
    lines = [
        ["配置编辑器使用说明"],
        [""],
        ["1. 本 Excel 是「编辑层」，改完保存后运行生成器写回 JSON，系统读取 JSON。"],
        ["   运行: python scripts/config_excel_to_json.py"],
        ["   统一设计(2026-09-28): 每张表 第1行=蓝色提示(这张表怎么用)、第2行=表头；【浅黄底=可编辑列】，浅灰底=只读/参考列。"],
        ["   销售归属表按母公司分块浅蓝斑马纹；时间配置/销售归属/字段映射 表头带筛选按钮、已冻结窗格；各表标签页颜色区分。"],
        ["   下拉只给取值有限的项（模式/动态策略/指标/部门/排序），客户名等自由文本自己填。"],
        [""],
        ["2. 「时间配置」sheet 列说明："],
        ["   - 配置名: 年度累计 / 月度数据 / 季度累计筛选 / 年基线数据（勿改）"],
        ["   - 模式: 手填日期（默认：自己填开始/结束日期）或 自动计算（按「动态策略」自动算，不用填日期）"],
        ["     · 兼容旧写法 static = 手填日期、dynamic = 自动计算，历史文件不用改"],
        ["   - 动态策略: 仅在「模式 = 自动计算」时选：最近完整月 / 最近完整季度"],
        ["     · 例：9 月跑批 → 最近完整月 = 8/1~8/31；最近完整季度 = 4/1~6/30"],
        ["     · 注意：它取的是「上一个完整周期」，不是当月；要统计「当月」请继续用「手填日期」"],
        ["   - 开始日期/结束日期: 格式 YYYY-MM-DD 或 YYYY-MM-DD HH:MM:SS（两列都显示到秒，如 2026-10-01 00:00:00）"],
        ["     · 结束只写到「日」时，系统按当天 23:59:59 处理（不会漏掉当天数据）"],
        ["   - 年份/月份范围: 仅年基线数据使用，如 2025 / 1-8"],
        ["     · 月份范围单元格是文本格式，请勿改成日期格式（否则 1-9 会被存成 2026/1/9，同步会报错）"],
        ["   - 说明: 填写配置用途（可选）"],
        [""],
        *RULE_NOTE_LINES,
        ["4. 「销售归属」sheet 列说明："],
        ["   - 母公司: 客户归组的母公司名（如 广州小鹏汽车科技有限公司）"],
        ["   - 子公司: 实际结算主体名（=母公司时表示本部）"],
        ["   - 指标: 收入 / 回款 / 收入/回款（斜杠分隔 = 两个指标同构时用后者，一行搞定）"],
        ["   - 部门: 检测 / 信息 / 能源 / 海外"],
        ["   - 销售: 负责销售的姓名"],
        ["   - 比例: 分配比例，默认 1（全额）；拆分时填小数如 0.3 或百分比 30%"],
        ["   - 说明: 备注（可选）"],
        *GROUP_NOTE_LINES,
        [""],
        ["5. 「KPI指标」sheet 列说明（Hero 圆环指标总数，万元）："],
        ["   - 页面: 年度达成 / 月度达成 / 季度达成"],
        ["   - 指标: 收入 / 回款"],
        ["   - 值(万元): 填数字 = 覆盖 KPI 卡片显示的指标总数（如 50000 = 5亿）"],
        ["     · 留空 = 不覆盖，KPI 卡片仍按指标文件合计（现状）"],
        ["     · 不影响表格/矩阵/部门卡里的明细指标（只改 Hero 圆环总数）"],
        ["   - 说明: 备注（可选）"],
        ["   - 注意: 数据总览页的年度 KPI 跟随「年度达成」的设置"],
        [""],
        ["6. 常见修改："],
        ["   - 换月份: 改「时间配置」的月度数据开始/结束日期"],
        ["   - 调整优先展示客户: 「展示规则」sheet → 第 4 列 配置项 = 优先展示 的那些行（一个客户一行，可下拉选）"],
        ["   - 显示全部客户: 把 配置项 = 优先展示 的所有行的「值」清空（或删行）"],
        ["   - 只显示指定客户: 配置项 = 客户筛选 的那些行，「值」填客户名（下拉选）；清空 = 不筛选"],
        ["   - 改销售归属: 在「销售归属」sheet 直接改对应行的 销售 / 比例"],
        ["   - 拆分多销售: 把一行复制成多行，各填不同销售和比例（合计需=1）"],
        [""],
        ["7. 「字段映射」sheet 列说明（清洗时的字段↔列名匹配，2026-09-17 新增）："],
        ["   - 数据源: 财务端 / 运营端（勿改）"],
        ["   - 来源: 收入 / 回款 / 广东公司 / 湖南公司 / 南方韶关（勿改）"],
        ["   - 标准字段: 系统内部字段名（勿改）"],
        ["   - 候选1..候选N: 按优先级从左到右尝试匹配源文件列名；留空 = 不参与"],
        ["     · 匹配规则是【全名精确匹配】（仅忽略首尾空白），不是关键词匹配"],
        ["     · 源文件列名改了，把新列名加到候选里保存即可，无需改 JSON/代码"],
        ["     · 同一字段内候选不能重复；每个字段至少保留 1 个候选"],
        ["   - 当前命中(只读): 上次清洗实际命中的列，写回时被忽略，仅供参考"],
        ["   - 说明: 备注（可选）"],
        ["   - 「排除标记」字段（仅财务端收入，回款无此字段）: 源表列名就叫「初始化」；值 = √ 的行属期初/历史单据迁移"],
        ["     清洗时整行剔除、不计入当月收入（2026-09-24 口径）；命中取值在 cleaning_config.json 的 排除标记值 里改"],
        ["   - 注意: 本 sheet 只控制「列映射」；文件名/引擎/金额乘数(金额除数)/过客户白名单等仍在 JSON 里改（金额单位换算见第 9 节）"],
        [""],
        ["8. 校验失败时生成器拒绝写回并提示原因，JSON 保持原样。"],
        ["   字段映射写回后 cleaning_config.json 会统一为标准 2 空格缩进格式（写前自动备份到 logs/）。"],
        [""],
        *UNIT_NOTE_LINES,
        [""],
        *NAME_NOTE_LINES,
        [""],
        *OUTPUT_NOTE_LINES,
        [""],
        *BIZ_NOTE_LINES,
    ]
    for row in lines:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 110
    ws2["A1"].font = Font(bold=True, size=14, color=COLOR_PRIMARY)
    ws2["A1"].fill = PatternFill("solid", fgColor=COLOR_BANNER_BG)
    ws2.row_dimensions[1].height = 28
    # 分节标题（"1." / "2." … 开头的行）加粗 + 浅蓝底，便于扫读
    for i in range(2, ws2.max_row + 1):
        v = ws2.cell(row=i, column=1).value
        if isinstance(v, str) and re.match(r"^\d+\.", v.strip()):
            c = ws2.cell(row=i, column=1)
            c.font = Font(bold=True, color=COLOR_PRIMARY, size=11.5)
            c.fill = PatternFill("solid", fgColor=COLOR_BANNER_BG)
            ws2.row_dimensions[i].height = 20
    finish(ws2, tab_color=TAB_COLORS["说明"])

    # ── Sheet 6: 字段映射（清洗列名映射的编辑层；「当前命中」由 logs/column_hits_*.txt 回填）──
    attach_mapping_sheet(
        wb,
        json.loads(CLEANING_CFG.read_text(encoding="utf-8")),
        read_latest_hits(),
    )

    # ── sheet 顺序整理：说明固定放最后 ──
    if "说明" in wb.sheetnames and wb.sheetnames[-1] != "说明":
        wb.move_sheet("说明", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("说明"))

    wb.save(EXCEL_PATH)
    try:
        shown = EXCEL_PATH.relative_to(BASE_DIR)
    except ValueError:  # 测试/自定义路径（不在项目内）
        shown = EXCEL_PATH
    print(f"✅ 已生成模板: {shown}")


def _note_blocks() -> list[list[list[str]]]:
    """说明书里由代码维护的分节（顺序 = 写入顺序；每项是一个完整分节的行列表）"""
    return [UNIT_NOTE_LINES, NAME_NOTE_LINES, OUTPUT_NOTE_LINES, BIZ_NOTE_LINES]


def _sync_note_block(ws, block: list[list[str]]) -> str:
    """让「说明」sheet 里某个分节与代码文案一致，返回 'updated' / 'appended' / 'same'

    - 找不到该分节标题 → 追加到末尾（前空一行）
    - 找到但正文有变（如「一次运行一个文件」改口径）→ **原地重写该分节**（不动其它行）
    """
    title = str(block[0][0]).strip()
    rows = {r: str(ws.cell(row=r, column=1).value).strip() for r in range(1, ws.max_row + 1)}
    start = next((r for r, v in rows.items() if v == title), None)

    if start is None:
        row = ws.max_row + 2
        for line in block:
            cell = ws.cell(row=row, column=1, value=line[0])
            if re.match(r"^\d+\.", str(line[0]).strip()):
                cell.font = Font(bold=True, color=COLOR_PRIMARY, size=11.5)
                cell.fill = PatternFill("solid", fgColor=COLOR_BANNER_BG)
            row += 1
        return "appended"

    # 该分节的行范围：标题 → 下一个分节标题（或表尾）之前，去掉尾部空行
    # （分节之间的空行是排版，不属于分节内容 —— 否则每次刷新都会删掉一行，越刷越短）
    end = ws.max_row
    for r in range(start + 1, ws.max_row + 1):
        if re.match(r"^\d+\.", rows.get(r, "")):
            end = r - 1
            break
    while end > start and not str(ws.cell(row=end, column=1).value or "").strip():
        end -= 1
    have = [str(ws.cell(row=r, column=1).value) for r in range(start, end + 1)]
    want = [line[0] for line in block]
    if have == want:
        return "same"

    need, cur = len(want), len(have)
    if need < cur:
        ws.delete_rows(start + need, cur - need)
    elif need > cur:
        ws.insert_rows(start + cur, need - cur)
    for i, text in enumerate(want):
        cell = ws.cell(row=start + i, column=1, value=text)
        if i == 0:
            cell.font = Font(bold=True, color=COLOR_PRIMARY, size=11.5)
            cell.fill = PatternFill("solid", fgColor=COLOR_BANNER_BG)
        else:
            cell.font = Font(size=11)
    return "updated"


def refresh_note_sheet() -> Path:
    """让「说明」sheet 与代码里的文案同步：缺的分节追加、改过的分节原地重写

    用途：说明书新增/修订章节（如第 11 节改「一次运行一个文件」）后不必跑 `--init`
    重建整本（重建会覆盖手工编辑）；本命令只碰代码维护的那几个分节，其它内容一行不动。
    """
    wb = openpyxl.load_workbook(EXCEL_PATH)
    if "说明" not in wb.sheetnames:
        raise ValueError("工作簿缺少「说明」sheet")
    ws = wb["说明"]
    changed = [_sync_note_block(ws, block) for block in _note_blocks()]
    if all(c == "same" for c in changed):
        return EXCEL_PATH
    wb.save(EXCEL_PATH)
    return EXCEL_PATH
