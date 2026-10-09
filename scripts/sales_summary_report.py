# -*- coding: utf-8 -*-
# Author: naohua168 <bai_bai168@qq.com>
"""销售完成度汇总表生成器

读取引擎层两个销售拆分表(实际) + 年度指标表(目标),生成
`output/销售完成度/销售完成度汇总_YYYYMMDD.xlsx`,2 个 sheet(销售收入 / 销售回款)。

每个销售一行;每个部门"实际 + 指标"相邻两列;单位统一为"元"
(引擎实际金额保留原始精度,指标表万元 ×10000);完成率百分比。

用法:
    python scripts/sales_summary_report.py                 # 生成到 output/销售完成度/
    python scripts/sales_summary_report.py --output=D:\a.xlsx   # 指定输出路径
"""
from __future__ import annotations

import argparse
import datetime
import json
import sys
from pathlib import Path

import pandas as pd

BASE_DIR = Path(__file__).parent.parent.resolve()
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from engine.core import business_rules as br  # noqa: E402
from engine.core.utils import rel_path  # noqa: E402  （提示路径统一相对项目根显示）

DEPTS = br.departments()        # 部门列表来自 业务规则.json（2026-09-29）
DEFAULT_TARGET_MULT = 10000.0   # 手工指标表默认单位 = 万元（可由配置 `金额乘数` 覆盖）

# 取不到配置时的回落路径（与 cleaning_config.json 的默认值一致）
_FALLBACK = {
    "销售收入": ("输出", "销售收入", "data/sheets/系统数据清理/销售收入/", "销售收入.xlsx"),
    "销售回款": ("输出", "销售回款", "data/sheets/系统数据清理/销售回款/", "销售回款.xlsx"),
    "年度收入总指标": ("手工指标表", "年度收入总指标", "data/sheets/手动维护/年度收入总指标/", "年度收入总指标.xlsx"),
    "年度回款总指标": ("手工指标表", "年度回款总指标", "data/sheets/手动维护/年度回款总指标/", "年度回款总指标.xlsx"),
}


def _load_cfg() -> dict:
    try:
        return json.loads((BASE_DIR / "config" / "清洗配置" / "cleaning_config.json").read_text(encoding="utf-8"))
    except Exception as e:
        _log(f"读取 cleaning_config.json 失败（改用默认路径）: {e}", "WARN")
        return {}


def _first_xlsx(folder_rel: str, fallback_name: str) -> Path:
    """目录下第一个 xlsx（兼容文件名变化）；目录不存在时回落固定文件名"""
    folder = BASE_DIR / folder_rel
    files = sorted(f for f in folder.glob("*.xlsx") if not f.name.startswith("~$")) if folder.is_dir() else []
    return files[0] if files else folder / fallback_name


def resolve_paths() -> dict[str, Path]:
    """输入文件路径：全部来自 cleaning_config.json（输出/手工指标表），不再写死（2026-09-29）"""
    cfg = _load_cfg()
    out: dict[str, Path] = {}
    for key, (section, sub, folder_rel, fallback_name) in _FALLBACK.items():
        rel = (cfg.get(section) or {}).get(sub, folder_rel)
        out[key] = _first_xlsx(str(rel), fallback_name)
    return out


def output_dir() -> Path:
    """输出目录：配置 `输出.销售完成度`（默认 output/销售完成度/）"""
    rel = (_load_cfg().get("输出") or {}).get("销售完成度", "output/销售完成度/")
    return BASE_DIR / str(rel)


def target_multiplier() -> float:
    """手工指标表金额 → 元的换算倍数（默认 10000 = 万元）

    支持 `数据源.手工指标表.金额乘数` 或顶层 `手工指标表.金额乘数`（数值；按 Sheet 的字典取首值）
    """
    cfg = _load_cfg()
    raw = ((cfg.get("数据源") or {}).get("手工指标表") or {}).get("金额乘数")
    if raw is None:
        raw = (cfg.get("手工指标表") or {}).get("金额乘数")
    if isinstance(raw, dict):
        raw = next(iter(raw.values()), None)
    try:
        return float(raw) if raw is not None else DEFAULT_TARGET_MULT
    except (TypeError, ValueError):
        return DEFAULT_TARGET_MULT


def _log(msg: str, level: str = "INFO"):
    pre = {"INFO": "  ", "WARN": "  ⚠️ ", "OK": "  ✅ "}
    print(f"{pre.get(level, '  ')}[汇总表] {msg}")


def load_sales_xlsx(path: Path) -> pd.DataFrame:
    """实际表: 排除"待确认"，金额保持原始(元)"""
    df = pd.read_excel(path)
    df["销售"] = df["销售"].astype(str).str.strip()
    df["事业部"] = df["事业部"].astype(str).str.strip()
    df["金额"] = pd.to_numeric(df["金额"], errors="coerce").fillna(0)
    return df[df["销售"] != "待确认"]


def load_tgt_xlsx(path: Path, mult: float = DEFAULT_TARGET_MULT) -> pd.DataFrame:
    """年度指标表: 万元 → 元（×mult，默认 10000；由配置 `金额乘数` 决定）"""
    df = pd.read_excel(path)
    df["销售"] = df["销售"].astype(str).str.strip()
    for c in DEPTS:
        if c not in df.columns:
            df[c] = 0
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0) * mult
    return df


def build_sheet(act_df: pd.DataFrame, tgt_df: pd.DataFrame) -> pd.DataFrame:
    """一个销售一行：每部门 实际+指标 相邻(元)，合计 + 完成率"""
    act = act_df.groupby(["销售", "事业部"], as_index=False)["金额"].sum()
    act_piv = act.pivot_table(index="销售", columns="事业部", values="金额",
                              aggfunc="sum", fill_value=0)
    for d in DEPTS:
        if d not in act_piv.columns:
            act_piv[d] = 0.0
    act_piv = act_piv[DEPTS]
    act_piv["实际合计"] = act_piv[DEPTS].sum(axis=1)

    tgt = tgt_df.groupby("销售", as_index=False)[DEPTS].sum().set_index("销售")
    tgt["指标合计"] = tgt[DEPTS].sum(axis=1)

    all_sales = sorted(set(tgt.index) | set(act_piv.index))
    rows = []
    for s in all_sales:
        a = act_piv.loc[s] if s in act_piv.index else pd.Series(0.0, index=DEPTS + ["实际合计"])
        t = tgt.loc[s] if s in tgt.index else pd.Series(0.0, index=DEPTS + ["指标合计"])
        row = {"销售": s}
        for d in DEPTS:
            row[f"{d}实际"] = round(float(a[d]), 6)   # 保留原数据精度,不取整
            row[f"{d}指标"] = round(float(t[d]), 6)
        row["实际合计"] = round(float(a["实际合计"]), 6)
        row["指标合计"] = round(float(t["指标合计"]), 6)
        row["完成率"] = (row["实际合计"] / row["指标合计"]) if row["指标合计"] else None
        rows.append(row)
    out = pd.DataFrame(rows)
    cols = ["销售"]
    for d in DEPTS:
        cols += [f"{d}实际", f"{d}指标"]
    cols += ["实际合计", "指标合计", "完成率"]
    out = out[cols]
    out = out.sort_values("指标合计", ascending=False, na_position="last").reset_index(drop=True)
    return out


def generate(output_path: Path | None = None) -> Path:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    paths = resolve_paths()
    mult = target_multiplier()
    act_inc = load_sales_xlsx(paths["销售收入"])
    act_pay = load_sales_xlsx(paths["销售回款"])
    tgt_inc = load_tgt_xlsx(paths["年度收入总指标"], mult)
    tgt_pay = load_tgt_xlsx(paths["年度回款总指标"], mult)

    inc_df = build_sheet(act_inc, tgt_inc)
    pay_df = build_sheet(act_pay, tgt_pay)

    if output_path is None:
        today = datetime.date.today().strftime("%Y%m%d")
        out_dir = output_dir()
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"销售完成度汇总_{today}.xlsx"

    try:
        writer_ctx = pd.ExcelWriter(output_path, engine="openpyxl")
    except PermissionError as e:  # 被 Excel/WPS 打开
        _log(f"❌ 写入失败：{rel_path(output_path)} 正被 Excel/WPS 等程序占用 → 请关闭该文件后重跑", "WARN")
        raise SystemExit(3) from e
    with writer_ctx as writer:
        inc_df.to_excel(writer, sheet_name="销售收入", index=False)
        pay_df.to_excel(writer, sheet_name="销售回款", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions["A"].width = 12
            for col in range(2, ws.max_column):
                ws.column_dimensions[get_column_letter(col)].width = 16
                for r in range(2, ws.max_row + 1):
                    c = ws.cell(row=r, column=col)
                    if isinstance(c.value, (int, float)):
                        c.number_format = "#,##0.######"
            last = ws.max_column
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=last)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0%"
                elif cell.value is None:
                    cell.value = "—"
                    cell.alignment = Alignment(horizontal="center")

    return output_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成销售完成度汇总表(2 sheet)")
    parser.add_argument("--output", "-o", default=None, help="输出文件路径(默认 output/销售完成度/销售完成度汇总_YYYYMMDD.xlsx)")
    args = parser.parse_args(argv)

    paths = resolve_paths()
    mult = target_multiplier()
    for label, p in [("销售收入", paths["销售收入"]), ("销售回款", paths["销售回款"]),
                     ("年度收入指标", paths["年度收入总指标"]), ("年度回款指标", paths["年度回款总指标"])]:
        shown = p.relative_to(BASE_DIR) if str(p).startswith(str(BASE_DIR)) else p
        if not p.exists():
            _log(f"缺少数据文件: {shown}（{label}）", "WARN")
            return 1
        _log(f"{label}: {shown}")

    _log(f"指标表金额换算: ×{mult:g} → 元")
    _log("开始生成")
    out = generate(Path(args.output) if args.output else None)
    _log(f"已生成: {rel_path(out)}", "OK")

    # 摘要
    import pandas as pd
    inc_df = pd.read_excel(out, sheet_name="销售收入")
    pay_df = pd.read_excel(out, sheet_name="销售回款")
    print(f"\n销售数: 收入 {len(inc_df)} / 回款 {len(pay_df)}")
    print(f"收入实际合计 {inc_df['实际合计'].sum():,.2f} 元 / 指标 {inc_df['指标合计'].sum():,.0f} 元")
    print(f"回款实际合计 {pay_df['实际合计'].sum():,.2f} 元 / 指标 {pay_df['指标合计'].sum():,.0f} 元")
    return 0


if __name__ == "__main__":
    sys.exit(main())
